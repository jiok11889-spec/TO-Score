# -*- coding: utf-8 -*-
"""엑셀 안전장치: 수정 직전 자동 백업 + 저장 후 바뀐 셀 목록 출력.

사용법: wb.save(path) 대신  safe_save(wb, path)
백업은 data/backups/TO Score_YYYYMMDD_HHMMSS.xlsx 로 남는다.
"""
import shutil
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


def _snapshot(path):
    """{(시트, 셀주소): 값} — 값만 비교한다(서식 제외)."""
    wb = load_workbook(path, data_only=False)
    snap = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    snap[(ws.title, cell.coordinate)] = cell.value
    return snap


def safe_save(wb, path, max_lines=50):
    path = Path(path)
    backup = None
    before = {}
    if path.exists():
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup_dir = path.parent / "backups"
        backup_dir.mkdir(exist_ok=True)
        backup = backup_dir / f"{path.stem}_{stamp}{path.suffix}"
        shutil.copy2(path, backup)
        before = _snapshot(backup)

    wb.save(path)

    after = _snapshot(path)
    changes = []
    for key in sorted(set(before) | set(after)):
        old, new = before.get(key), after.get(key)
        if old != new:
            sheet, coord = key
            changes.append(f"  [{sheet}] {coord}: {old!r} -> {new!r}")

    if backup:
        print(f"백업 생성: {backup.name}")
    print(f"변경 셀 {len(changes)}개:")
    for line in changes[:max_lines]:
        print(line)
    if len(changes) > max_lines:
        print(f"  … 외 {len(changes) - max_lines}개")
    return backup
